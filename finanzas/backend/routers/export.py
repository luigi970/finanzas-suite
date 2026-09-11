import io
from datetime import datetime

import httpx
from fastapi import APIRouter
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database import get_db
from routers.agent import build_price_context

router = APIRouter(prefix="/api/export", tags=["export"])

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _style_header(ws, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _autosize(ws, ncols: int):
    for col in range(1, ncols + 1):
        letter = get_column_letter(col)
        max_len = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=0)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 45)


@router.get("/portfolio")
async def export_portfolio():
    """Excel con 3 hojas: Portfolio (posiciones + P&L, misma valuación que usa el
    agente IA — ver build_price_context), Cuentas (resumen por cuenta) y Movimientos
    (historial completo). No recalcula precios/P&L con lógica propia: reusa
    build_price_context tal cual para no repetir meses de bugs de precisión ya
    corregidos ahí (CEDEAR vs acción real, promedio ponderado multi-cuenta, etc.)."""
    conn = get_db()
    accounts = conn.execute("SELECT * FROM accounts WHERE active = 1 ORDER BY name").fetchall()
    positions = [
        dict(p) for p in conn.execute(
            "SELECT p.*, a.name as account_name FROM positions p JOIN accounts a ON p.account_id = a.id"
        ).fetchall()
    ]
    transactions = conn.execute(
        "SELECT t.*, a.name as account_name FROM transactions t JOIN accounts a ON t.account_id = a.id "
        "ORDER BY t.date DESC"
    ).fetchall()
    conn.close()

    async with httpx.AsyncClient(timeout=30) as client:
        _, rows = await build_price_context(positions, client)

    wb = Workbook()

    # ── Portfolio ──────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Portfolio"
    headers = ["Cuenta", "Activo", "Tipo", "Cantidad", "Precio mercado USD", "Valor USD",
               "Precio prom. compra USD", "P&L no realizado USD", "P&L no realizado %", "Notas"]
    ws.append(headers)
    for r in rows:
        ws.append([
            r["account_name"], r["asset"], r["asset_type"], round(r["quantity"], 6),
            round(r["market_price_usd"], 4) if r["market_price_usd"] is not None else None,
            round(r["value_usd"], 2) if r["value_usd"] is not None else None,
            round(r["avg_price_usd"], 4) if r["avg_price_usd"] is not None else None,
            round(r["unrealized_pnl_usd"], 2) if r["unrealized_pnl_usd"] is not None else None,
            round(r["unrealized_pnl_pct"], 2) if r["unrealized_pnl_pct"] is not None else None,
            r["notes"] or "",
        ])
    # Totales sumados en Python, nunca dejados para que Excel/el lector los infiera mal
    # (mismo criterio que TOTAL CARTERA / P&L NO REALIZADO TOTAL en build_price_context).
    total_value = sum(r["value_usd"] for r in rows if r["value_usd"] is not None)
    total_pnl = sum(r["unrealized_pnl_usd"] for r in rows if r["unrealized_pnl_usd"] is not None)
    ws.append([])
    ws.append(["TOTAL", "", "", "", "", round(total_value, 2), "", round(total_pnl, 2), "", ""])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    _style_header(ws, len(headers))
    _autosize(ws, len(headers))

    # ── Cuentas ────────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Cuentas")
    ws2.append(["Cuenta", "Tipo", "Valor total USD"])
    value_by_account = {}
    for r in rows:
        if r["value_usd"] is not None:
            value_by_account[r["account_name"]] = value_by_account.get(r["account_name"], 0.0) + r["value_usd"]
    for a in accounts:
        ws2.append([a["name"], a["type"], round(value_by_account.get(a["name"], 0.0), 2)])
    ws2.append([])
    ws2.append(["TOTAL", "", round(sum(value_by_account.values()), 2)])
    for cell in ws2[ws2.max_row]:
        cell.font = Font(bold=True)
    _style_header(ws2, 3)
    _autosize(ws2, 3)

    # ── Movimientos ────────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Movimientos")
    headers3 = ["Fecha", "Cuenta", "Descripción", "Monto", "Moneda", "Tipo", "Categoría",
                "Precio unit.", "P&L realizado", "Comisión", "Moneda comisión", "Fuente"]
    ws3.append(headers3)
    for t in transactions:
        ws3.append([
            t["date"], t["account_name"], t["description"] or "", t["amount"], t["currency"],
            t["type"], t["category"] or "", t["unit_price"], t["realized_pnl"],
            t["fee"], t["fee_currency"] or "", t["source"] or "",
        ])
    _style_header(ws3, len(headers3))
    _autosize(ws3, len(headers3))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"finanzas_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
